"""Tests for the Modulo Rosa Società (roster imports) feature."""
import io
import os
import pytest
import requests
import openpyxl

def _load_env():
    if os.environ.get("REACT_APP_BACKEND_URL"):
        return os.environ["REACT_APP_BACKEND_URL"]
    for line in open("/app/frontend/.env"):
        if line.startswith("REACT_APP_BACKEND_URL="):
            return line.split("=", 1)[1].strip()
    raise RuntimeError("REACT_APP_BACKEND_URL not set")

BASE = _load_env().rstrip("/") + "/api"
TID = "6a9b5b045d9e0985643d0a9f"
ROMA_TEAM = "6a9b5b045d9e0985643d0aaf"
OTHER_TEAM = "6a9b5b045d9e0985643d0ab3"  # Sporting Eur 2014 (different club)


def _login(email, password):
    s = requests.Session()
    r = s.post(f"{BASE}/auth/login", json={"email": email, "password": password}, timeout=15)
    assert r.status_code == 200, r.text
    return s


@pytest.fixture(scope="module")
def admin():
    return _login("castellani.giampaolo@gmail.com", os.environ["ADMIN_PASSWORD"])


@pytest.fixture(scope="module")
def club():
    return _login("societa@fsl.demo", os.environ["QA_PASSWORD"])


def _template_bytes(club_sess):
    r = club_sess.get(f"{BASE}/tournaments/{TID}/roster-imports/template", params={"team_id": ROMA_TEAM}, timeout=15)
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers.get("content-type", "")
    return r.content


def _fill(template_bytes, players, coach="Mister Test", contact="Test Ref", phone="+39 000", team_name=None):
    wb = openpyxl.load_workbook(io.BytesIO(template_bytes))
    ws = wb.active
    if team_name is not None:
        ws["B5"] = team_name
    ws["B6"] = coach
    ws["B7"] = contact
    ws["E7"] = phone
    for idx, p in enumerate(players):
        r = 11 + idx
        ws.cell(r, 2).value = p[0]  # first
        ws.cell(r, 3).value = p[1]  # last
        ws.cell(r, 4).value = p[2]  # role
        ws.cell(r, 5).value = p[3]  # shirt
        ws.cell(r, 6).value = p[4]  # birth
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def _clean_submitted(admin, club_id=None):
    """Reject any pending submitted imports so tests can create new ones."""
    r = admin.get(f"{BASE}/tournaments/{TID}/roster-imports", params={"status": "submitted"}, timeout=15)
    for i in r.json():
        admin.post(f"{BASE}/tournaments/{TID}/roster-imports/{i['id']}/reject", json={"note": "cleanup TEST"}, timeout=15)


# ---------------- Template ----------------

def test_template_download_prefill(club):
    data = _template_bytes(club)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.active
    assert "Roma Nord" in str(ws["B5"].value)


# ---------------- Submit ----------------

def test_submit_non_xlsx_rejected(club):
    files = {"file": ("bad.txt", b"not xlsx", "text/plain")}
    r = club.post(f"{BASE}/tournaments/{TID}/roster-imports", data={"team_id": ROMA_TEAM}, files=files, timeout=15)
    assert r.status_code == 400


def test_submit_empty_roster_rejected(club):
    tpl = _template_bytes(club)
    xlsx = _fill(tpl, [])
    files = {"file": ("empty.xlsx", xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    r = club.post(f"{BASE}/tournaments/{TID}/roster-imports", data={"team_id": ROMA_TEAM}, files=files, timeout=15)
    assert r.status_code == 400


def test_submit_other_club_forbidden(club):
    tpl = _template_bytes(club)
    xlsx = _fill(tpl, [("Mario", "Test", "Portiere", 1, "01/01/2014")])
    files = {"file": ("r.xlsx", xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    r = club.post(f"{BASE}/tournaments/{TID}/roster-imports", data={"team_id": OTHER_TEAM}, files=files, timeout=15)
    assert r.status_code == 403


def test_submit_with_errors_and_conflict(admin, club):
    _clean_submitted(admin)
    tpl = _template_bytes(club)
    xlsx = _fill(
        tpl,
        [
            ("TESTAlfa", "Rossi", "Portiere", 1, "01/01/2014"),          # ok
            ("TESTBeta", "Bianchi", "SuperEroe", 7, "05/06/2014"),        # unknown role
            ("TESTGamma", "Neri", "Attaccante", 1, "10/03/2014"),          # duplicate shirt 1
            ("TESTDelta", "Gialli", "Difensore", 4, "99/99/9999"),          # invalid date
            ("", "Blu", "Centrocampista", 5, "01/01/2014"),                 # missing name
        ],
        team_name="Roma Nord Test Squad",
    )
    files = {"file": ("r.xlsx", xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    r = club.post(f"{BASE}/tournaments/{TID}/roster-imports", data={"team_id": ROMA_TEAM}, files=files, timeout=15)
    assert r.status_code == 201, r.text
    body = r.json()
    assert body["team_name"] == "Roma Nord Test Squad"
    assert body["coach"] == "Mister Test"
    assert body["phone"] == "+39 000"
    rows = body["rows"]
    assert len(rows) == 5
    errs = {row["n"]: " | ".join(row["errors"]) for row in rows}
    assert "Ruolo non riconosciuto" in errs[2]
    assert "duplicato" in errs[3]
    assert "Data" in errs[4] or "data" in errs[4].lower()
    assert "Nome" in errs[5] or "cognome" in errs[5].lower()

    # Second submission while one is submitted -> 409
    xlsx2 = _fill(tpl, [("Uno", "Due", "Portiere", 2, "01/01/2014")])
    files = {"file": ("r2.xlsx", xlsx2, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    r2 = club.post(f"{BASE}/tournaments/{TID}/roster-imports", data={"team_id": ROMA_TEAM}, files=files, timeout=15)
    assert r2.status_code == 409

    # Save id for later tests
    pytest.import_id = body["id"]


def test_list_club_scope_and_staff_filter(admin, club):
    r_club = club.get(f"{BASE}/tournaments/{TID}/roster-imports", timeout=15)
    assert r_club.status_code == 200
    club_items = r_club.json()
    assert all(i["club_name"].startswith("Roma Nord") or "Roma Nord" in i["club_name"] for i in club_items)
    assert any(i["id"] == pytest.import_id for i in club_items)

    r_all = admin.get(f"{BASE}/tournaments/{TID}/roster-imports", timeout=15)
    assert r_all.status_code == 200
    assert len(r_all.json()) >= len(club_items)

    r_sub = admin.get(f"{BASE}/tournaments/{TID}/roster-imports", params={"status": "submitted"}, timeout=15)
    assert r_sub.status_code == 200
    assert all(i["status"] == "submitted" for i in r_sub.json())


# ---------------- Approve / Reject ----------------

def test_club_manager_cannot_approve(club):
    r = club.post(f"{BASE}/tournaments/{TID}/roster-imports/{pytest.import_id}/approve", json={"mode": "merge"}, timeout=15)
    assert r.status_code == 403


def test_approve_missing_role_rejected(admin):
    bad_rows = [{"n": 1, "first_name": "TESTAlfa", "last_name": "Rossi", "role": "", "shirt_number": 1, "birth_year": 2014}]
    r = admin.post(f"{BASE}/tournaments/{TID}/roster-imports/{pytest.import_id}/approve",
                   json={"mode": "merge", "rows": bad_rows}, timeout=15)
    assert r.status_code == 400


def test_approve_merge_success(admin):
    # Fix the erroneous rows and only include valid ones
    fixed_rows = [
        {"n": 1, "first_name": "TESTAlfa", "last_name": "Rossi", "role": "Portiere", "shirt_number": 21, "birth_year": 2014},
        {"n": 2, "first_name": "TESTBeta", "last_name": "Bianchi", "role": "Attaccante", "shirt_number": 22, "birth_year": 2014},
    ]
    r = admin.post(f"{BASE}/tournaments/{TID}/roster-imports/{pytest.import_id}/approve",
                   json={"mode": "merge", "rows": fixed_rows, "note": "OK test"}, timeout=15)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["status"] == "approved"
    assert body["created"] + body["updated"] == 2

    # Approving again -> 409
    r2 = admin.post(f"{BASE}/tournaments/{TID}/roster-imports/{pytest.import_id}/approve",
                    json={"mode": "merge", "rows": fixed_rows}, timeout=15)
    assert r2.status_code == 409

    # Verify players created in team
    rp = admin.get(f"{BASE}/tournaments/{TID}/players", params={"team_id": ROMA_TEAM}, timeout=15)
    assert rp.status_code == 200
    names = {(p["first_name"], p["last_name"]) for p in rp.json()}
    assert ("Testalfa", "Rossi") in names or ("TESTAlfa", "Rossi") in names or any(n[0].lower() == "testalfa" for n in names)


def test_reject_without_note(admin, club):
    # Create a new pending submission first
    tpl = _template_bytes(club)
    xlsx = _fill(tpl, [("TESTReject", "Uno", "Portiere", 30, "01/01/2014")])
    files = {"file": ("rj.xlsx", xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    rsub = club.post(f"{BASE}/tournaments/{TID}/roster-imports", data={"team_id": ROMA_TEAM}, files=files, timeout=15)
    assert rsub.status_code == 201
    iid = rsub.json()["id"]

    r_no_note = admin.post(f"{BASE}/tournaments/{TID}/roster-imports/{iid}/reject", json={"note": ""}, timeout=15)
    assert r_no_note.status_code == 400

    r_ok = admin.post(f"{BASE}/tournaments/{TID}/roster-imports/{iid}/reject", json={"note": "Motivo test"}, timeout=15)
    assert r_ok.status_code == 200
    assert r_ok.json()["status"] == "rejected"


def test_club_notifications_created(club):
    r = club.get(f"{BASE}/tournaments/{TID}/notifications", timeout=15)
    assert r.status_code == 200
    items = r.json().get("items", [])
    titles = " | ".join(n.get("title", "") for n in items).lower()
    assert "caricata" in titles or "respinto" in titles


# ---------------- Replace mode ----------------

def test_replace_mode_sets_missing_inactive(admin, club):
    # First get current active players in team
    rp = admin.get(f"{BASE}/tournaments/{TID}/players", params={"team_id": ROMA_TEAM}, timeout=15)
    existing_active = [p for p in rp.json() if p["status"] == "active"]
    assert len(existing_active) >= 2

    # Submit a roster containing ONLY one existing player + a new one
    keep = existing_active[0]
    tpl = _template_bytes(club)
    xlsx = _fill(tpl, [
        (keep["first_name"], keep["last_name"], keep["role"], keep.get("shirt_number") or 50, "01/01/2014"),
        ("TESTReplace", "Nuovo", "Difensore", 88, "01/01/2014"),
    ])
    files = {"file": ("rep.xlsx", xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    rs = club.post(f"{BASE}/tournaments/{TID}/roster-imports", data={"team_id": ROMA_TEAM}, files=files, timeout=15)
    assert rs.status_code == 201
    iid = rs.json()["id"]

    ra = admin.post(f"{BASE}/tournaments/{TID}/roster-imports/{iid}/approve",
                    json={"mode": "replace", "rows": rs.json()["rows"]}, timeout=15)
    assert ra.status_code == 200, ra.text

    # Player who was excluded should now be inactive
    rp2 = admin.get(f"{BASE}/tournaments/{TID}/players", params={"team_id": ROMA_TEAM}, timeout=15)
    active_names_after = {(p["first_name"].lower(), p["last_name"].lower()) for p in rp2.json() if p["status"] == "active"}
    excluded = existing_active[1]
    assert (excluded["first_name"].lower(), excluded["last_name"].lower()) not in active_names_after
