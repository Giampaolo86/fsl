import io
from datetime import date, datetime
from typing import Optional

from fastapi import APIRouter, Depends, File, Form, Response, UploadFile
from pydantic import BaseModel

from ..core.deps import CurrentUser, get_current_user, require_tournament
from ..core.errors import bad_request, conflict, forbidden, not_found
from ..services.linkcodes import new_code
from ..models.domain import Player, RosterImport
from ..repositories.registry import scoped
from ..services import audit
from .club_extras import _store, notify

router = APIRouter(prefix="/tournaments/{tournament_id}/roster-imports", tags=["roster-imports"])
STAFF = {"super_admin", "director", "secretary"}
ROLES = {"portiere": "Portiere", "difensore": "Difensore", "centrocampista": "Centrocampista", "esterno": "Esterno", "attaccante": "Attaccante", "por": "Portiere", "dif": "Difensore", "cen": "Centrocampista", "att": "Attaccante"}
FIRST_ROW = 11


def _txt(v) -> str:
    return "" if v is None else str(v).strip()


def _birth(v):
    if v is None or v == "":
        return None, None
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d"), v.year
    s = str(v).strip()
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d.%m.%Y"):
        try:
            d = datetime.strptime(s, fmt)
            return d.strftime("%Y-%m-%d"), d.year
        except ValueError:
            continue
    if s.isdigit() and len(s) == 4:
        return None, int(s)
    return s, None


def parse_workbook(data: bytes) -> dict:
    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    except Exception:  # noqa: BLE001
        raise bad_request("File non valido: usa il modulo Excel (.xlsx) scaricato dalla piattaforma")
    ws = wb.active
    head = {"team_name": _txt(ws["B5"].value), "coach": _txt(ws["B6"].value), "contact": _txt(ws["B7"].value), "phone": _txt(ws["E7"].value)}
    rows, numbers, year_now = [], {}, date.today().year
    for r in range(FIRST_ROW, ws.max_row + 1):
        first, last, role, num, born = (ws.cell(r, c).value for c in range(2, 7))
        if not any(_txt(x) for x in (first, last, role, num, born)):
            continue
        errors = []
        first, last = _txt(first).title(), _txt(last).title()
        if not first or not last:
            errors.append("Nome o cognome mancante")
        role_n = ROLES.get(_txt(role).lower(), "")
        if not role_n:
            errors.append("Ruolo non riconosciuto")
        shirt = None
        try:
            shirt = int(float(num)) if _txt(num) else None
            if shirt is not None and not 1 <= shirt <= 99:
                errors.append("Numero maglia fuori range (1-99)")
        except ValueError:
            errors.append("Numero maglia non numerico")
        if shirt is not None:
            if shirt in numbers:
                errors.append(f"Numero {shirt} duplicato con riga {numbers[shirt]}")
            numbers[shirt] = r - FIRST_ROW + 1
        birth_date, birth_year = _birth(born)
        if born not in (None, "") and birth_year is None:
            errors.append("Data di nascita non valida (gg/mm/aaaa)")
        if birth_year and not (year_now - 20 <= birth_year <= year_now - 4):
            errors.append("Anno di nascita non plausibile")
        rows.append({"n": r - FIRST_ROW + 1, "first_name": first, "last_name": last, "role": role_n or _txt(role), "shirt_number": shirt, "birth_date": birth_date, "birth_year": birth_year, "errors": errors})
    return {**head, "rows": rows}


async def _out(t_id: str, items: list[RosterImport]) -> list[dict]:
    clubs = {c.id: c for c in await scoped("clubs", t_id).list()}
    teams = {tm.id: tm for tm in await scoped("teams", t_id).list()}
    out = []
    for i in items:
        d = i.public()
        d["club_name"] = clubs[i.club_id].name if i.club_id in clubs else ""
        d["team_label"] = teams[i.team_id].name if i.team_id in teams else ""
        d["team_status"] = teams[i.team_id].status if i.team_id in teams else ""
        d["team_category"] = teams[i.team_id].category if i.team_id in teams else ""
        d["error_count"] = sum(1 for r in i.rows if r.get("errors"))
        out.append(d)
    return out


@router.get("/template")
async def template(tournament_id: str, team_id: Optional[str] = None, club_id: Optional[str] = None, user: CurrentUser = Depends(get_current_user)):
    from ..services.roster_template import build_template

    t, role = await require_tournament(tournament_id, user)
    tm = await scoped("teams", tournament_id).get(team_id) if team_id else None
    if role == "club_manager" and tm and tm.club_id != user.club_in(tournament_id):
        raise forbidden("Puoi scaricare solo il modulo della tua società")
    club = await scoped("clubs", tournament_id).get(tm.club_id if tm else (club_id or user.club_in(tournament_id) or "")) if (tm or club_id or user.club_in(tournament_id)) else None
    comp_name = ""
    if tm and tm.competition_id:
        comp = await scoped("competitions", tournament_id).get(tm.competition_id)
        comp_name = comp.name if comp else ""
    name = "Modulo_Rosa"
    if club:
        name = f"Modulo_Rosa_{(club.short_name or club.name).replace(' ', '_')}{('_' + tm.category) if tm and tm.category else ''}"
    data = build_template(t, club, tm, comp_name)
    return Response(content=data, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="FSL_{name}.xlsx"'})


@router.post("", status_code=201)
async def submit(tournament_id: str, team_id: str = Form(...), file: UploadFile = File(...), user: CurrentUser = Depends(get_current_user)):
    t, role = await require_tournament(tournament_id, user, roles=STAFF | {"club_manager"}, writable=True)
    tm = await scoped("teams", tournament_id).get(team_id)
    if not tm:
        raise not_found("Squadra")
    if role == "club_manager" and tm.club_id != user.club_in(tournament_id):
        raise forbidden("Puoi caricare solo la rosa della tua società")
    if not (file.filename or "").lower().endswith(".xlsx"):
        raise bad_request("Carica il modulo in formato .xlsx")
    data = await file.read()
    if len(data) > 5 * 1024 * 1024:
        raise bad_request("File troppo grande (max 5 MB)")
    parsed = parse_workbook(data)
    if not parsed["rows"]:
        raise bad_request("Il modulo non contiene giocatori: compila la tabella ROSA GIOCATORI")
    repo = scoped("roster_imports", tournament_id)
    if await repo.find_one({"team_id": team_id, "status": "submitted"}):
        raise conflict("C'è già un modulo in attesa di conferma per questa squadra")
    m = await _store(tournament_id, data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", file.filename, user.id, tm.club_id)
    imp = await repo.insert(RosterImport(tournament_id=tournament_id, club_id=tm.club_id, team_id=team_id, media_id=m.id, file_url=f"/api/media/{m.id}", filename=file.filename, submitted_by=user.id, **parsed), user.id)
    await audit.record(user, "roster_import.submit", "roster_import", imp.id, tournament_id, after={"team_id": team_id, "rows": len(parsed["rows"]), "errors": sum(1 for r in parsed["rows"] if r["errors"])})
    if role == "club_manager":
        from ..services import staff_notify

        club = await scoped("clubs", tournament_id).get(tm.club_id)
        pend = " · gruppo da confermare" if tm.status == "pending" else ""
        await staff_notify.notify_staff(tournament_id, "roster", f"Rosa caricata: {club.name if club else ''} · {tm.category}", f"{tm.name} · {len(parsed['rows'])} giocatori{pend} · {t.name}. Rivedi e conferma da Rose.", f"/admin/t/{tournament_id}/rose", f"roster:{imp.id}")
    return (await _out(tournament_id, [imp]))[0]


@router.get("")
async def list_imports(tournament_id: str, status: Optional[str] = None, user: CurrentUser = Depends(get_current_user)):
    t, role = await require_tournament(tournament_id, user, roles=STAFF | {"club_manager"})
    f = {}
    if role == "club_manager":
        f["club_id"] = user.club_in(tournament_id)
    if status:
        f["status"] = {"$in": status.split(",")}
    return await _out(tournament_id, await scoped("roster_imports", tournament_id).list(f, sort=[("created_at", -1)], limit=200))


class ApproveIn(BaseModel):
    mode: str = "merge"
    rows: Optional[list[dict]] = None
    note: str = ""


class AddPlayerIn(BaseModel):
    team_id: str
    first_name: str
    last_name: str
    role: str = ""
    shirt_number: Optional[int] = None
    birth_year: Optional[int] = None
    note: str = ""


@router.post("/request-player", status_code=201)
async def request_player(tournament_id: str, body: AddPlayerIn, user: CurrentUser = Depends(get_current_user)):
    t, role = await require_tournament(tournament_id, user, roles=STAFF | {"club_manager"}, writable=True)
    tm = await scoped("teams", tournament_id).get(body.team_id)
    if not tm:
        raise not_found("Squadra")
    if role == "club_manager" and tm.club_id != user.club_in(tournament_id):
        raise forbidden("Puoi richiedere giocatori solo per la tua società")
    first, last = body.first_name.strip(), body.last_name.strip()
    if len(first) < 2 or len(last) < 2:
        raise bad_request("Inserisci nome e cognome")
    errors = []
    role_n = ROLES.get(body.role.strip().lower()) if body.role else ""
    if body.role and not role_n:
        errors.append("Ruolo non riconosciuto")
    if body.shirt_number is not None and not (1 <= body.shirt_number <= 99):
        errors.append("Numero maglia fuori range (1-99)")
    existing = await scoped("players", tournament_id).find_one({"team_id": tm.id, "shirt_number": body.shirt_number}) if body.shirt_number is not None else None
    if existing:
        errors.append(f"Numero {body.shirt_number} già assegnato a {existing.first_name} {existing.last_name}")
    row = {"n": 1, "first_name": first, "last_name": last, "role": role_n or body.role.strip(), "shirt_number": body.shirt_number, "birth_date": "", "birth_year": body.birth_year, "errors": errors}
    repo = scoped("roster_imports", tournament_id)
    imp = await repo.insert(RosterImport(tournament_id=tournament_id, club_id=tm.club_id, team_id=tm.id, filename="Richiesta aggiunta giocatore", team_name=tm.name, rows=[row], note=body.note.strip()[:300], submitted_by=user.id), user.id)
    await audit.record(user, "roster_import.request_player", "roster_import", imp.id, tournament_id, after={"team_id": tm.id, "player": f"{first} {last}"})
    return (await _out(tournament_id, [imp]))[0]


@router.post("/{import_id}/approve")
async def approve(tournament_id: str, import_id: str, body: ApproveIn, user: CurrentUser = Depends(get_current_user)):
    await require_tournament(tournament_id, user, roles=STAFF, writable=True)
    repo = scoped("roster_imports", tournament_id)
    imp = await repo.get(import_id)
    if not imp:
        raise not_found("Modulo")
    if imp.status != "submitted":
        raise conflict("Modulo già gestito")
    if body.mode not in ("merge", "replace"):
        raise bad_request("Modalità non valida")
    rows = body.rows if body.rows is not None else imp.rows
    clean = []
    for r in rows:
        first, last = _txt(r.get("first_name")).title(), _txt(r.get("last_name")).title()
        role_n = ROLES.get(_txt(r.get("role")).lower(), "")
        if not first or not last or not role_n:
            raise bad_request(f"Riga {r.get('n')}: nome, cognome e ruolo sono obbligatori")
        shirt = r.get("shirt_number")
        shirt = int(shirt) if shirt not in (None, "") else None
        by = r.get("birth_year")
        by = int(by) if by not in (None, "") else None
        clean.append({"first_name": first, "last_name": last, "role": role_n, "shirt_number": shirt, "birth_year": by})
    nums = [c["shirt_number"] for c in clean if c["shirt_number"] is not None]
    if len(nums) != len(set(nums)):
        raise bad_request("Numeri di maglia duplicati nel modulo")
    players = scoped("players", tournament_id)
    existing = await players.list({"team_id": imp.team_id}, limit=500)
    created = updated = 0
    if body.mode == "replace":
        keys = {(c["first_name"].lower(), c["last_name"].lower()) for c in clean}
        for p in existing:
            if (p.first_name.lower(), p.last_name.lower()) not in keys and p.status != "inactive":
                await players.update(p.id, {"status": "inactive"}, user.id)
    by_key = {(p.first_name.lower(), p.last_name.lower()): p for p in existing}
    for c in clean:
        p = by_key.get((c["first_name"].lower(), c["last_name"].lower()))
        if p:
            await players.update(p.id, {"role": c["role"], "shirt_number": c["shirt_number"], "birth_year": c["birth_year"] or p.birth_year, "status": "active"}, user.id)
            updated += 1
        else:
            await players.insert(Player(link_code=new_code(), tournament_id=tournament_id, club_id=imp.club_id, team_id=imp.team_id, **c), user.id)
            created += 1
    imp2 = await repo.update(imp.id, {"status": "approved", "rows": rows, "note": body.note, "reviewed_by": user.id, "imported_count": created + updated}, user.id)
    from ..services import staff_notify

    await staff_notify.settle(f"roster:{imp.id}")
    teams_repo = scoped("teams", tournament_id)
    pending_tm = await teams_repo.get(imp.team_id)
    if pending_tm and pending_tm.status == "pending":
        await teams_repo.update(pending_tm.id, {"status": "active"}, user.id)
        await staff_notify.settle(f"group:{pending_tm.id}")
        await audit.record(user, "team.confirm", "team", pending_tm.id, tournament_id, after={"via": "roster_import"})
    await audit.record(user, "roster_import.approve", "roster_import", imp.id, tournament_id, after={"mode": body.mode, "created": created, "updated": updated})
    tm = await scoped("teams", tournament_id).get(imp.team_id)
    await notify(tournament_id, imp.club_id, "roster", f"Rosa «{tm.name if tm else imp.team_name}» caricata", f"{created} nuovi giocatori, {updated} aggiornati." + (f" Nota: {body.note}" if body.note else ""), "/societa/rose", f"rosterimp:{imp.id}:approved")
    d = (await _out(tournament_id, [imp2]))[0]
    d.update({"created": created, "updated": updated})
    return d


@router.post("/{import_id}/reject")
async def reject(tournament_id: str, import_id: str, body: dict, user: CurrentUser = Depends(get_current_user)):
    await require_tournament(tournament_id, user, roles=STAFF, writable=True)
    repo = scoped("roster_imports", tournament_id)
    imp = await repo.get(import_id)
    if not imp:
        raise not_found("Modulo")
    if imp.status != "submitted":
        raise conflict("Modulo già gestito")
    note = _txt(body.get("note"))
    if not note:
        raise bad_request("Indica il motivo del rifiuto")
    imp2 = await repo.update(imp.id, {"status": "rejected", "note": note, "reviewed_by": user.id}, user.id)
    from ..services import staff_notify

    await staff_notify.settle(f"roster:{imp.id}")
    await audit.record(user, "roster_import.reject", "roster_import", imp.id, tournament_id, reason=note)
    await notify(tournament_id, imp.club_id, "roster", "Modulo rosa respinto", note, "/societa/rose", f"rosterimp:{imp.id}:rejected")
    return (await _out(tournament_id, [imp2]))[0]
