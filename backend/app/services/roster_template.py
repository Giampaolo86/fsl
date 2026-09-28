import io
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

NAVY, NAVY2, GOLD, WHITE, SLATE, LIGHT = "0B1F3A", "0F2A4F", "F4AE2B", "FFFFFF", "6B7A90", "F3F6FB"
LOGO = Path(__file__).resolve().parents[3] / "frontend" / "public" / "brand" / "logo.png"
ROLES = ["Portiere", "Difensore", "Centrocampista", "Esterno", "Attaccante"]
FIRST_ROW, ROWS = 11, 40
COLS = ["N.", "NOME", "COGNOME", "RUOLO", "N° MAGLIA", "DATA DI NASCITA (gg/mm/aaaa)"]
WIDTHS = [6, 24, 24, 22, 13, 24]


def _fill(hex_):
    return PatternFill("solid", fgColor=hex_)


def _font(size=11, bold=False, color=NAVY, italic=False, name="Arial"):
    return Font(name=name, size=size, bold=bold, color=color, italic=italic)


def build_template(tournament=None, club=None, team=None, competition: str = "") -> bytes:
    """Modulo rosa in stile FSL: intestazione navy/oro con logo, dati precompilati, 40 righe validate, stampa A4."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Rosa"
    ws.sheet_view.showGridLines = False
    for i, w in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    thin = Side(style="thin", color="C9D2E0")
    gold_side = Side(style="medium", color=GOLD)
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)

    # Intestazione (righe 1-2): logo + titolo su fondo navy
    ws.row_dimensions[1].height = 34
    ws.row_dimensions[2].height = 34
    for r in (1, 2):
        for c in range(1, 7):
            ws.cell(r, c).fill = _fill(NAVY)
    ws.merge_cells("B1:F1")
    ws.merge_cells("B2:F2")
    ws["B1"] = "FUTURE STARS LEAGUE"
    ws["B1"].font = _font(20, True, WHITE)
    ws["B1"].alignment = Alignment(horizontal="left", vertical="bottom", indent=1)
    ws["B2"] = "MODULO ROSA SOCIETÀ"
    ws["B2"].font = _font(14, True, GOLD)
    ws["B2"].alignment = Alignment(horizontal="left", vertical="top", indent=1)
    if LOGO.exists():
        img = XLImage(str(LOGO))
        img.width = img.height = 80
        ws.add_image(img, "A1")
    # Striscia oro (riga 3)
    ws.merge_cells("A3:F3")
    ws.row_dimensions[3].height = 20
    ws["A3"] = f"La Serie A del futuro · {tournament.name if tournament else 'Torneo'}{(' · ' + tournament.season_label) if tournament is not None and getattr(tournament, 'season_label', '') else ''}"
    ws["A3"].fill = _fill(GOLD)
    ws["A3"].font = _font(10, True, NAVY)
    ws["A3"].alignment = center
    # Istruzioni (riga 4)
    ws.merge_cells("A4:F4")
    ws.row_dimensions[4].height = 58
    ws["A4"] = ("Compila una riga per ogni atleta: nome, cognome, ruolo (menu a tendina), numero di maglia (1-99, senza doppioni) e data di nascita nel formato gg/mm/aaaa. "
                "Non modificare la struttura del foglio. Salva in formato .xlsx e ricaricalo nell'Area Società → Rose (o consegnalo alla segreteria FSL): "
                "la rosa viene creata automaticamente dopo la conferma dell'organizzazione.")
    ws["A4"].font = _font(9, color=SLATE, italic=True)
    ws["A4"].alignment = left
    # Dati squadra (righe 5-8)
    labels = {5: "SQUADRA", 6: "ALLENATORE", 7: "REFERENTE SOCIETÀ"}
    for r, lab in labels.items():
        ws.row_dimensions[r].height = 22
        ws.cell(r, 1, lab).font = _font(9, True, WHITE)
        ws.cell(r, 1).fill = _fill(NAVY2)
        ws.cell(r, 1).alignment = center
    ws.merge_cells("B5:F5")
    ws.merge_cells("B6:F6")
    ws.merge_cells("B7:C7")
    ws.merge_cells("E7:F7")
    ws["D7"] = "TELEFONO"
    ws["D7"].font = _font(9, True, WHITE)
    ws["D7"].fill = _fill(NAVY2)
    ws["D7"].alignment = center
    ws["B5"] = " · ".join(x for x in [club.name if club else "", team.name if team else ""] if x)
    for ref in ("B5", "B6", "B7", "E7"):
        ws[ref].font = _font(11, ref == "B5", NAVY)
        ws[ref].alignment = left
        ws[ref].border = Border(bottom=gold_side)
    ws.row_dimensions[8].height = 22
    ws.cell(8, 1, "CATEGORIA").font = _font(9, True, WHITE)
    ws.cell(8, 1).fill = _fill(NAVY2)
    ws.cell(8, 1).alignment = center
    ws["B8"] = getattr(team, "category", "") or ""
    ws["B8"].font = _font(11, True, NAVY)
    ws["B8"].alignment = left
    ws["C8"] = "COMPETIZIONE"
    ws["C8"].font = _font(9, True, WHITE)
    ws["C8"].fill = _fill(NAVY2)
    ws["C8"].alignment = center
    ws.merge_cells("D8:F8")
    ws["D8"] = competition or ""
    ws["D8"].font = _font(11, color=NAVY)
    ws["D8"].alignment = left
    # Titolo tabella (riga 9) e intestazioni (riga 10)
    ws.merge_cells("A9:F9")
    ws.row_dimensions[9].height = 26
    ws["A9"] = "ROSA GIOCATORI"
    ws["A9"].fill = _fill(GOLD)
    ws["A9"].font = _font(12, True, NAVY)
    ws["A9"].alignment = center
    ws.row_dimensions[10].height = 30
    for c, h in enumerate(COLS, start=1):
        cell = ws.cell(10, c, h)
        cell.fill = _fill(NAVY)
        cell.font = _font(9, True, WHITE)
        cell.alignment = center
        cell.border = box
    for i in range(ROWS):
        r = FIRST_ROW + i
        ws.row_dimensions[r].height = 20
        for c in range(1, 7):
            cell = ws.cell(r, c)
            cell.border = box
            cell.font = _font(10, c == 1, SLATE if c == 1 else NAVY)
            cell.alignment = center if c in (1, 4, 5, 6) else left
            if i % 2 == 1:
                cell.fill = _fill(LIGHT)
        ws.cell(r, 1, i + 1)
        ws.cell(r, 6).number_format = "DD/MM/YYYY"
    last = FIRST_ROW + ROWS - 1
    dv_role = DataValidation(type="list", formula1='"' + ",".join(ROLES) + '"', allow_blank=True, showErrorMessage=True, errorTitle="Ruolo", error="Scegli un ruolo dal menu: " + ", ".join(ROLES))
    dv_num = DataValidation(type="whole", operator="between", formula1="1", formula2="99", allow_blank=True, showErrorMessage=True, errorTitle="Numero di maglia", error="Inserisci un numero intero da 1 a 99")
    dv_date = DataValidation(type="date", operator="between", formula1="DATE(2000,1,1)", formula2=f"DATE({date.today().year},12,31)", allow_blank=True, showErrorMessage=True, errorTitle="Data di nascita", error="Inserisci la data nel formato gg/mm/aaaa (es. 14/03/2014)")
    for dv, col in ((dv_role, "D"), (dv_num, "E"), (dv_date, "F")):
        ws.add_data_validation(dv)
        dv.add(f"{col}{FIRST_ROW}:{col}{last}")
    # Nota finale
    ws.merge_cells(f"A{last + 2}:F{last + 2}")
    ws.row_dimensions[last + 2].height = 40
    ws[f"A{last + 2}"] = "Con la consegna del modulo la società conferma di aver verificato i dati anagrafici dei tesserati e di aver raccolto i consensi privacy/immagine delle famiglie. I dati restano riservati: sul portale compaiono nome e foto solo con consenso attivo."
    ws[f"A{last + 2}"].font = _font(8, color=SLATE, italic=True)
    ws[f"A{last + 2}"].alignment = left
    # Stampa A4
    ws.print_area = f"A1:F{last + 2}"
    ws.print_title_rows = "10:10"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.oddFooter.center.text = "Future Stars League · La Serie A del futuro · Modulo rosa società"
    ws.oddFooter.center.size = 8
    ws.freeze_panes = f"A{FIRST_ROW}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
